"""Read-only serial-to-customer evidence from explicitly authorized SharePoint files.

This module is NOT a Microsoft Graph client. A governed, identity-bound SharePoint
reader must separately authorize and provide every versioned document. The
ChatGPT delegated connector is never a Jason service credential. There is no
filesystem scan, arbitrary URL fetch, API call, write, or customer-identity guess.
"""
from __future__ import annotations

import base64
import binascii
import hashlib
import hmac
import io
import re
from zipfile import BadZipFile, ZipFile
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Iterable, Mapping

SERIAL_RE = re.compile(r"(?<![A-Z0-9])(?=[A-Z0-9-]{8,24}\b)(?=[A-Z0-9-]*[A-Z])(?=[A-Z0-9-]*\d)[A-Z0-9-]{8,24}(?![A-Z0-9])", re.I)
LABEL_RE = re.compile(r"^(?:serial(?:\s*(?:no\.?|number|#))?|s\s*/?\s*n\s*(?:#|no\.?)?)\s*[:#]?\s*(.*)$", re.I)
RECEIPT_COMPANY_RE = re.compile(r"^(?:co\.?\s*name|company(?:\s+name)?|customer(?:\s+name)?)\s*:?\s*$", re.I)
SERIAL_VALID_RE = re.compile(r"^[A-Z0-9-]{8,24}$")


@dataclass(frozen=True, slots=True)
class ApprovedSharePointWorkbook:
    """Evidence envelope supplied only after Jason source and access authorization."""

    tenant_id: str
    site_id: str
    drive_id: str
    item_id: str
    etag: str
    web_url: str
    filename: str
    modified_at: datetime
    content: bytes

    def __post_init__(self) -> None:
        if not all(
            isinstance(v, str) and v.strip()
            for v in (self.tenant_id, self.site_id, self.drive_id, self.item_id,
                      self.etag, self.web_url, self.filename)
        ):
            raise ValueError("versioned source identity is required")
        if not self.filename.casefold().endswith(".xlsx"):
            raise ValueError("only non-macro Excel documents are supported")
        if self.modified_at.tzinfo is None:
            raise ValueError("source modification time must include timezone")
        if not isinstance(self.content, bytes) or len(self.content) > 10 * 1024 * 1024:
            raise ValueError("document must be bounded binary evidence")
        if not re.search(r"copier|meter|log|receipt", self.filename, re.I):
            raise ValueError("document is not an approved copier-log or receipt class")
        if not self.content.startswith(b"PK"):
            raise ValueError("approved Excel document is not a ZIP container")


def document_from_governed_sharepoint_reads(
    *,
    tenant_id: str,
    site_id: str,
    drive_id: str,
    item_id: str,
    metadata_before: object,
    content_result: object,
    metadata_after: object,
) -> ApprovedSharePointWorkbook:
    """Bridge the existing governed SharePoint reader's verified output format.

    Callers MUST have separately authenticated the Jason requester, tenant,
    site/library/item release policy, and all three read operations. Requiring
    identical before/after eTags prevents accepting bytes from a file that
    changed during the read. This adapter never fetches network resources.
    """
    def evidence(result: object, *, operation: str) -> Mapping[str, object]:
        if getattr(result, "provider", None) != "microsoft_graph":
            raise PermissionError("source must be the governed Microsoft Graph reader")
        if getattr(result, "capability", None) != operation:
            raise PermissionError("unexpected governed SharePoint operation")
        data = getattr(result, "data", None)
        if not isinstance(data, Mapping):
            raise ValueError("invalid governed SharePoint result data")
        return data

    before = evidence(metadata_before, operation="knowledge.document.metadata.read").get("item")
    after = evidence(metadata_after, operation="knowledge.document.metadata.read").get("item")
    content = evidence(content_result, operation="knowledge.document.content.read")
    if not isinstance(before, Mapping) or not isinstance(after, Mapping):
        raise ValueError("versioned source metadata is required")
    required = ("id", "name", "eTag", "size", "lastModifiedDateTime", "webUrl", "site_id", "drive_id")
    if any(not before.get(key) and before.get(key) != 0 for key in required):
        raise ValueError("governed source metadata is incomplete")
    if any(before.get(key) != after.get(key) for key in required):
        raise ValueError("source document changed during read")
    if before["id"] != item_id or before["site_id"] != site_id or before["drive_id"] != drive_id:
        raise PermissionError("source identity does not match authorized selectors")
    if not isinstance(before["size"], int) or isinstance(before["size"], bool):
        raise ValueError("source document size is invalid")
    if not 0 < before["size"] <= 10 * 1024 * 1024:
        raise ValueError("source document exceeds safe size")
    if not str(before["name"]).casefold().endswith(".xlsx"):
        raise ValueError("source must be an approved xlsx document")
    file_meta = before.get("file")
    if not isinstance(file_meta, Mapping) or file_meta.get("mimeType") != (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        raise ValueError("source media type is not approved")
    b64 = content.get("content_b64")
    if not isinstance(b64, str) or len(b64) > 15 * 1024 * 1024:
        raise ValueError("bounded source content is unavailable")
    if content.get("encoding") != "base64":
        raise ValueError("governed document encoding is invalid")
    try:
        raw = base64.b64decode(b64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError("invalid base64 document evidence") from exc
    if len(raw) != before["size"] or content.get("byte_length") != len(raw):
        raise ValueError("source byte size differs from authorized metadata")
    digest = hashlib.sha256(raw).hexdigest()
    if not isinstance(content.get("sha256"), str) or not hmac.compare_digest(
        digest, content["sha256"]
    ):
        raise ValueError("source digest validation failed")
    modified_at = datetime.fromisoformat(
        str(before["lastModifiedDateTime"]).replace("Z", "+00:00")
    )
    return ApprovedSharePointWorkbook(
        tenant_id=tenant_id,
        site_id=site_id,
        drive_id=drive_id,
        item_id=item_id,
        etag=str(before["eTag"]),
        filename=str(before["name"]),
        web_url=str(before["webUrl"]),
        modified_at=modified_at,
        content=raw,
    )


@dataclass(frozen=True, slots=True)
class CustomerClaim:
    serial_number: str
    customer_name: str
    source_item_id: str
    source_drive_id: str
    source_etag: str
    source_sha256: str
    source_url: str
    source_sheet: str
    modified_at: str


@dataclass(frozen=True, slots=True)
class CustomerMatch:
    status: str  # verified | conflict | not_found
    customer_name: str | None
    evidence: tuple[CustomerClaim, ...] = ()


def _text(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def _customer_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _serial(value: object) -> str | None:
    word = _text(value).upper()
    if not SERIAL_VALID_RE.fullmatch(word) or not any(c.isalpha() for c in word):
        return None
    if not any(c.isdigit() for c in word):
        return None
    return word


def _first_customer(rows: list[tuple[object, ...]]) -> str | None:
    """Accept only a single explicit sheet-wide customer.

    Multi-customer tables or stacked receipts require a different bounded row
    contract. Mapping every serial on such a sheet to its first company is
    dangerous, so those sheets deliberately produce no customer assertions.
    """
    candidates: dict[str, str] = {}
    for row_index, row in enumerate(rows[:5000]):
        labels = [_text(v).casefold().strip(":") for v in row[:32]]
        if "company name" in labels and "serial number" in labels:
            return None  # Multi-customer row table; not a single-customer log.
        for col, raw in enumerate(row[:24]):
            if not RECEIPT_COMPANY_RE.fullmatch(_text(raw)):
                continue
            adjacent = tuple(_text(x) for x in row[col + 1:col + 3])
            candidate = next((v for v in adjacent if v), "")
            if not candidate and row_index + 1 < len(rows):
                below = rows[row_index + 1]
                if col < len(below):
                    candidate = _text(below[col])
            if candidate and candidate.casefold() not in {
                "company", "customer", "unnamed", "e-copiers", "ecopiers solutions",
            }:
                if len(candidate) <= 160 and not SERIAL_RE.fullmatch(candidate):
                    candidates[_customer_key(candidate)] = candidate
    return next(iter(candidates.values())) if len(candidates) == 1 else None


def _customer_from_named_log(filename: str, rows: list[tuple[object, ...]]) -> str | None:
    """Corroborate a customer-named meter log with a matching body heading.

    A document title alone is not identity evidence. Require the same distinctive
    first two name tokens in a top-of-sheet header and an exact serial elsewhere.
    Generic multi-customer logs do not qualify.
    """
    match = re.fullmatch(r"(.+?)\s+Log\s*\([^)]*\)\.xlsx", filename, re.I)
    if match is None:
        return None
    candidate = _text(match.group(1))
    parts = _customer_key(candidate).split()
    if len(parts) < 2 or len(parts[1]) < 4:
        return None
    witness = " ".join(parts[:2])
    headings = " ".join(_customer_key(_text(value)) for row in rows[:5] for value in row[:16])
    if not re.search(r"\b" + re.escape(witness) + r"\b", headings):
        return None
    if len(candidate) > 160 or "ecopier" in _customer_key(candidate):
        return None
    return candidate


def _serials_on_sheet(rows: list[tuple[object, ...]]) -> tuple[str, ...]:
    matches: set[str] = set()
    for row in rows[:5000]:
        for column, value in enumerate(row[:32]):
            text = _text(value)
            label = LABEL_RE.fullmatch(text)
            if label is None:
                # A labeled machine row may contain "ID#... SN: SERIAL" inline.
                inline = re.search(
                    r"(?:\bserial\s*(?:#|no\.?|number)?|\bs\s*/?\s*n\s*#?)\s*:?\s+([A-Z0-9-]{8,24})",
                    text, re.I,
                )
                if inline:
                    serial = _serial(inline.group(1))
                    if serial:
                        matches.add(serial)
                continue
            tail = label.group(1)
            serial = _serial(tail)
            if serial:
                matches.add(serial)
                continue
            for next_value in row[column + 1:column + 4]:
                serial = _serial(next_value)
                if serial:
                    matches.add(serial)
                    break
    return tuple(sorted(matches))


def _parse_workbook(doc: ApprovedSharePointWorkbook) -> tuple[CustomerClaim, ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("governed SharePoint workbook extraction requires openpyxl") from exc

    try:
        with ZipFile(io.BytesIO(doc.content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 60 * 1024 * 1024:
                raise ValueError("compressed SharePoint workbook exceeds extraction bounds")
    except BadZipFile as exc:
        raise ValueError("invalid SharePoint workbook archive") from exc

    workbook = load_workbook(
        io.BytesIO(doc.content), read_only=True, data_only=True, keep_links=False,
    )
    try:
        if len(workbook.worksheets) > 200:
            raise ValueError("workbook has too many sheets")
        document_hash = hashlib.sha256(doc.content).hexdigest()
        claims: list[CustomerClaim] = []
        inspected_rows = 0
        for sheet in workbook.worksheets:
            if not isinstance(sheet.title, str) or len(sheet.title) > 100:
                raise ValueError("invalid source worksheet")
            # Do not trust declared worksheet dimensions for a complete read.
            sheet.reset_dimensions()
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if index > 5000:
                    raise ValueError("worksheet row limit exceeded")
                if len(row) > 64:
                    raise ValueError("worksheet column limit exceeded")
                inspected_rows += 1
                if inspected_rows > 25000:
                    raise ValueError("total workbook row limit exceeded")
                rows.append(tuple(row))
            customer = _first_customer(rows)
            if customer is None:
                customer = _customer_from_named_log(doc.filename, rows)
            if customer is None:
                continue
            for serial in _serials_on_sheet(rows):
                claims.append(CustomerClaim(
                    serial_number=serial,
                    customer_name=customer,
                    source_item_id=doc.item_id,
                    source_drive_id=doc.drive_id,
                    source_etag=doc.etag,
                    source_sha256=document_hash,
                    source_url=doc.web_url,
                    source_sheet=sheet.title,
                    modified_at=doc.modified_at.astimezone(timezone.utc).isoformat(),
                ))
            if len(claims) > 2500:
                raise ValueError("workbook customer mappings exceed bounded limit")
        return tuple(claims)
    finally:
        workbook.close()


@dataclass(frozen=True, slots=True)
class ApprovedCustomerEvidenceIndex:
    """Bounded immutable cross-document index, no mutation or provider transport."""

    _claims_by_serial: Mapping[str, tuple[CustomerClaim, ...]]

    @classmethod
    def from_documents(
        cls,
        *,
        documents: Iterable[ApprovedSharePointWorkbook],
        tenant_id: str,
        allowed_site_ids: frozenset[str],
        allowed_drive_ids: frozenset[str],
        now: datetime,
        max_source_age_days: int = 400,
    ) -> "ApprovedCustomerEvidenceIndex":
        if not tenant_id or not allowed_site_ids or not allowed_drive_ids:
            raise PermissionError("approved tenant, site and drive boundaries are required")
        if now.tzinfo is None or not 1 <= max_source_age_days <= 730:
            raise ValueError("valid review time and freshness window required")
        batches = tuple(documents)
        if len(batches) > 40:
            raise ValueError("too many supplemental source documents")
        registry: dict[str, list[CustomerClaim]] = {}
        seen: set[tuple[str, str]] = set()
        for doc in batches:
            if (doc.tenant_id != tenant_id or doc.site_id not in allowed_site_ids
                    or doc.drive_id not in allowed_drive_ids):
                raise PermissionError("document source is outside approved AOT scope")
            key = (doc.drive_id, doc.item_id)
            if key in seen:
                raise ValueError("duplicate versioned source identity")
            seen.add(key)
            if doc.modified_at > now + timedelta(minutes=5):
                raise ValueError("source modified_at cannot be in the future")
            if now - doc.modified_at > timedelta(days=max_source_age_days):
                continue  # A stale source never becomes positive identity evidence.
            for claim in _parse_workbook(doc):
                registry.setdefault(claim.serial_number, []).append(claim)
        return cls({
            serial: tuple(values)
            for serial, values in registry.items()
        })

    def match(self, *, serial_number: str, kfs_customer: str | None = None) -> CustomerMatch:
        serial = _serial(serial_number)
        if serial is None:
            return CustomerMatch("not_found", None)
        evidence = self._claims_by_serial.get(serial, ())
        if not evidence:
            return CustomerMatch("not_found", None)
        names = {_customer_key(claim.customer_name) for claim in evidence}
        if len(names) != 1:
            return CustomerMatch("conflict", None, evidence)
        customer = evidence[0].customer_name
        if kfs_customer and _customer_key(kfs_customer) != _customer_key(customer):
            return CustomerMatch("conflict", None, evidence)
        return CustomerMatch("verified", customer, evidence)
