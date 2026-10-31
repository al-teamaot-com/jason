from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import pytest

from jason_runtime.kfs_meter_history import KfsMeterHistoryStore


class FakeCursor:
    def __init__(self, rows):
        self.rows = rows
        self.sql = ""
        self.params = None

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def execute(self, sql, params):
        self.sql = sql
        self.params = params

    def fetchall(self):
        return self.rows


class FakeConnection:
    def __init__(self, cursor):
        self.cur = cursor

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def cursor(self):
        return self.cur


def store_with_rows(monkeypatch, rows, *, expected_count=None, errors=0,
                    snapshot_at=None, snapshot_id="run-1", inventory_count=None):
    store = KfsMeterHistoryStore("unused", 5432, "unused", "reader", Path("/unused"))
    device_count = len({row[0] for row in rows})
    if expected_count is None:
        expected_count = device_count
    if inventory_count is None:
        inventory_count = device_count
    if snapshot_at is None and rows:
        snapshot_at = max(row[5] for row in rows if row[5] is not None)
    expanded = [
        (*row, snapshot_id, snapshot_at, expected_count, errors, inventory_count)
        for row in rows
    ]
    cur = FakeCursor(expanded)
    monkeypatch.setattr(KfsMeterHistoryStore, "_connect", lambda self: FakeConnection(cur))
    return store, cur


def test_fleet_ranking_distinguishes_missing_boundary_and_sorts(monkeypatch):
    from datetime import timedelta
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = []
    for device, total, bw, color in (("a", 40, 20, 20), ("b", 100, 80, 20)):
        for counter, delta in (("total", total), ("blackWhite", bw), ("color", color)):
            rows.append((device, counter, 100, start, 100 + delta, end,
                         "Customer", device.upper(), device))
    rows.append(("c", "total", None, None, 150, end, "Customer", "C", "c"))
    rows.append(("c", "blackWhite", None, None, 100, end, "Customer", "C", "c"))
    rows.append(("c", "color", None, None, 50, end, "Customer", "C", "c"))
    store, cursor = store_with_rows(monkeypatch, rows)
    result = store.fleet_rank(arguments={"limit": 2}, start=start, end=end)
    assert [item["device_id"] for item in result["top_devices"]] == ["b", "a"]
    assert result["top_devices"][0]["total_pages"] == 100
    assert result["top_devices"][0]["black_white_pages"] == 80
    assert result["top_devices"][0]["color_pages"] == 20
    assert result["exception_count"] == 1
    assert result["complete"] is False
    assert result["fleet_devices_in_history"] == 3
    assert len(cursor.params) == 11


def test_fleet_ranking_rejects_unbounded_limit(monkeypatch):
    store, _ = store_with_rows(monkeypatch, [])
    with pytest.raises(ValueError, match="limit"):
        store.fleet_rank(arguments={"limit": 9999}, start=datetime.now(timezone.utc),
                         end=datetime.now(timezone.utc))


def test_future_month_fleet_ranking_rejected_before_db(monkeypatch):
    from jason_runtime.kfs_meter_history import GovernedKfsMeterHistoryInvoker
    from orchestrator.print_capability_catalog import PRINT_FLEET_USAGE_RANK, KYOCERA_KFS_PROVIDER
    from types import SimpleNamespace
    store, cursor = store_with_rows(monkeypatch, [])
    invoker = GovernedKfsMeterHistoryInvoker(store)
    request = SimpleNamespace(permission_mode="observe", organization_id="aot",
                              client_id=None, capability_name=PRINT_FLEET_USAGE_RANK,
                              arguments={"month": "2099-01"})
    resolution = SimpleNamespace(selected_provider_id=KYOCERA_KFS_PROVIDER,
                                 capability_name=PRINT_FLEET_USAGE_RANK)
    with pytest.raises(ValueError, match="must have started"):
        invoker.invoke(request=request, resolution=resolution)
    assert cursor.sql == ""


def test_fleet_ranking_435_devices_matches_authoritative_snapshot(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = []
    for index in range(435):
        device = f"dev-{index:03d}"
        total = index + 1
        for counter, delta in (("total", total), ("blackWhite", total), ("color", 0)):
            rows.append((device, counter, 100, start, 100 + delta, end,
                         "Client", device, device))
    store, _ = store_with_rows(monkeypatch, rows)
    result = store.fleet_rank(arguments={"limit": 5}, start=start, end=end)
    assert result["fleet_devices_in_history"] == 435
    assert result["eligible_devices"] == 435
    assert result["exception_count"] == 0
    assert [x["device_id"] for x in result["top_devices"]] == [
        "dev-434", "dev-433", "dev-432", "dev-431", "dev-430",
    ]
    assert result["fleet_devices_reported"] == 435
    assert result["fleet_devices_in_inventory"] == 435
    assert result["coverage_verified"] is True
    assert result["complete"] is True


def test_fleet_ranking_rejects_identical_boundary_timestamps(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("dev-1", counter, 100, start, 100 + delta, start, "Client", "ID1", "S1")
        for counter, delta in (("total", 20), ("blackWhite", 10), ("color", 10))
    ]
    store, _ = store_with_rows(monkeypatch, rows)
    result = store.fleet_rank(arguments={}, start=start, end=end)
    assert result["eligible_devices"] == 0
    assert result["exception_count"] == 1
    assert result["top_devices"] == []

def test_fleet_ranking_fails_closed_on_inventory_count_mismatch(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("a", counter, 10, start, 10 + value, end, "Client", "A", "A")
        for counter, value in (("total", 7), ("blackWhite", 6), ("color", 1))
    ]
    store, _ = store_with_rows(monkeypatch, rows, expected_count=435)
    result = store.fleet_rank(arguments={}, start=start, end=end)
    assert result["fleet_devices_reported"] == 435
    assert result["fleet_devices_in_inventory"] == 1
    assert result["complete"] is False
    assert any("differs" in warning for warning in result["warnings"])


def test_fleet_ranking_fails_closed_on_stale_inventory(monkeypatch):
    from datetime import timedelta
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("a", counter, 10, start, 10 + value, end, "Client", "A", "A")
        for counter, value in (("total", 7), ("blackWhite", 6), ("color", 1))
    ]
    store, _ = store_with_rows(
        monkeypatch, rows, snapshot_at=end - timedelta(hours=40),
    )
    result = store.fleet_rank(arguments={}, start=start, end=end)
    assert result["complete"] is False
    assert result["coverage_verified"] is False
    assert any("stale" in warning for warning in result["warnings"])


def test_fleet_ranking_excludes_missing_inventory_devices_from_complete_claim(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("a", counter, 10, start, 10 + value, end, "Client", "A", "A")
        for counter, value in (("total", 7), ("blackWhite", 6), ("color", 1))
    ]
    store, _ = store_with_rows(monkeypatch, rows, inventory_count=2)
    result = store.fleet_rank(arguments={}, start=start, end=end)
    assert result["complete"] is False
    assert any("missing" in warning for warning in result["warnings"])


def test_monochrome_device_without_color_counter_uses_exact_zero(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("mono-1", "total", 100, start, 145, end, "Mono Client", "M-1", "S1"),
        ("mono-1", "blackWhite", 100, start, 145, end, "Mono Client", "M-1", "S1"),
        ("mono-1", "color", None, None, None, None, "Mono Client", "M-1", "S1"),
    ]
    store, cur = store_with_rows(monkeypatch, rows)
    report = store.fleet_rank(arguments={"limit": 5}, start=start, end=end)
    assert report["fleet_devices_reported"] == 1
    assert report["eligible_devices"] == 1
    assert report["exception_count"] == 0
    assert report["coverage_verified"] is True
    assert report["complete"] is True
    assert report["top_devices"][0]["total_pages"] == 45
    assert report["top_devices"][0]["black_white_pages"] == 45
    assert report["top_devices"][0]["color_pages"] == 0
    assert report["top_devices"][0]["color_counter_status"] == "zero_verified_from_total_and_bw"
    assert report["color_zero_verified_count"] == 1
    assert report["customer_mapping_missing_count"] == 0
    assert "color: missing meter boundary" not in report["top_devices"][0]["warnings"]
    assert "LEFT JOIN kfs_devices" in cur.sql


def test_missing_color_counter_must_not_invent_nonzero_delta(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("color-1", "total", 100, start, 150, end, "Client", "C-1", "S1"),
        ("color-1", "blackWhite", 100, start, 120, end, "Client", "C-1", "S1"),
        ("color-1", "color", None, None, None, None, "Client", "C-1", "S1"),
    ]
    store, _ = store_with_rows(monkeypatch, rows)
    report = store.fleet_rank(arguments={}, start=start, end=end)
    assert report["eligible_devices"] == 0
    assert report["exception_count"] == 1
    assert report["complete"] is False


def test_missing_customer_is_explicit_and_prevents_complete_claim(monkeypatch):
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("unmapped-1", name, 10, start, 10 + value, end, "", "E-1", "S1")
        for name, value in (("total", 7), ("blackWhite", 7), ("color", 0))
    ]
    store, _ = store_with_rows(monkeypatch, rows)
    report = store.fleet_rank(arguments={}, start=start, end=end)
    assert report["eligible_devices"] == 1
    assert report["exception_count"] == 0
    assert report["coverage_verified"] is True
    assert report["complete"] is False
    assert report["top_devices"][0]["customer"] == "Unknown (not mapped in KFS)"
    assert report["top_devices"][0]["customer_mapping_verified"] is False
    assert report["customer_mapping_missing_count"] == 1
    assert any("customer assignment" in message for message in report["warnings"])


def test_fleet_report_enriches_missing_customer_from_authorized_sharepoint_evidence(monkeypatch):
    from dataclasses import replace
    from datetime import timedelta
    from io import BytesIO
    from openpyxl import Workbook
    from jason_runtime.kfs_customer_document_evidence import (
        ApprovedCustomerEvidenceIndex, ApprovedSharePointWorkbook,
    )
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = []
    for device_id, serial, total, bw, color in (
        ("a", "1G55503749", 125, 125, 0),
        ("b", "RRG2301432", 50, 50, 0),
    ):
        for counter, delta in (("total", total), ("blackWhite", bw), ("color", color)):
            rows.append((device_id, counter, 100, start, 100 + delta, end,
                         "", "", serial))
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "CO. NAME:"
    ws["B1"] = "Marcella Operator LLC"
    ws["A2"] = "S/N:"
    ws["B2"] = "1G55503749"
    ws2 = wb.create_sheet("St Matthews")
    ws2["A1"] = "CO. NAME:"
    ws2["B1"] = "St. Matthews Catholic School"
    ws2["A2"] = "S/N:"
    ws2["B2"] = "RRG2301432"
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    doc = ApprovedSharePointWorkbook(
        tenant_id="aot", site_id="approved", drive_id="admin-documents",
        item_id="doc-1", etag="etag-1",
        web_url="https://aot.sharepoint.com/copiers/document",
        filename="AOT Copier Delivery Receipts.xlsx",
        modified_at=end - timedelta(days=2), content=buf.getvalue(),
    )
    evidence = ApprovedCustomerEvidenceIndex.from_documents(
        documents=(doc,), tenant_id="aot",
        allowed_site_ids=frozenset({"approved"}),
        allowed_drive_ids=frozenset({"admin-documents"}), now=end,
    )
    store, _ = store_with_rows(monkeypatch, rows)
    store = replace(store, customer_evidence_index=evidence)
    result = store.fleet_rank(arguments={"limit": 2}, start=start, end=end)
    assert result["fleet_devices_reported"] == 2
    assert result["supplemental_mapping_count"] == 2
    assert result["customer_mapping_missing_count"] == 0
    assert result["customer_mapping_conflict_count"] == 0
    assert result["complete"] is True
    assert [item["customer"] for item in result["top_devices"]] == [
        "Marcella Operator LLC", "St. Matthews Catholic School",
    ]
    assert all(item["customer_mapping_source"] == "aot_sharepoint" for item in result["top_devices"])
    assert all(item["customer_evidence"][0]["item_id"] == "doc-1" for item in result["top_devices"])


def test_fleet_report_marks_conflicting_document_and_kfs_customer(monkeypatch):
    from dataclasses import replace
    from datetime import timedelta
    from io import BytesIO
    from openpyxl import Workbook
    from jason_runtime.kfs_customer_document_evidence import (
        ApprovedCustomerEvidenceIndex, ApprovedSharePointWorkbook,
    )
    start = datetime(2026, 10, 1, tzinfo=timezone.utc)
    end = datetime(2026, 10, 30, tzinfo=timezone.utc)
    rows = [
        ("a", counter, 100, start, 100 + value, end, "KFS Customer", "A", "1G55503749")
        for counter, value in (("total", 12), ("blackWhite", 12), ("color", 0))
    ]
    wb = Workbook()
    wb.active.append(("CO. NAME:", "Other Customer"))
    wb.active.append(("S/N:", "1G55503749"))
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    doc = ApprovedSharePointWorkbook(
        tenant_id="aot", site_id="approved", drive_id="admin-documents",
        item_id="doc-2", etag="etag-2",
        web_url="https://aot.sharepoint.com/copiers/document-2",
        filename="AOT Copier Delivery Receipts.xlsx",
        modified_at=end - timedelta(days=2), content=buf.getvalue(),
    )
    index = ApprovedCustomerEvidenceIndex.from_documents(
        documents=(doc,), tenant_id="aot",
        allowed_site_ids=frozenset({"approved"}),
        allowed_drive_ids=frozenset({"admin-documents"}), now=end,
    )
    store, _ = store_with_rows(monkeypatch, rows)
    report = replace(store, customer_evidence_index=index).fleet_rank(
        arguments={}, start=start, end=end,
    )
    assert report["customer_mapping_conflict_count"] == 1
    assert report["complete"] is False
    assert report["top_devices"][0]["customer"] == "Conflict (requires customer review)"
    assert report["top_devices"][0]["customer_mapping_verified"] is False
