"""Negative and positive acceptance for approved AOT copier workbook evidence."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import Workbook

from jason_runtime.kfs_customer_document_evidence import (
    ApprovedCustomerEvidenceIndex,
    ApprovedSharePointWorkbook,
)

NOW = datetime(2026, 10, 30, tzinfo=timezone.utc)
TENANT = "aot-tenant"
SITE = "approved-admin-site"
DRIVE = "approved-admin-document-library"


def workbook_bytes(sheets: list[tuple[str, list[tuple]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets:
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def source(*, sheet, name="Copier log.xlsx", item="document-1",
           site=SITE, drive=DRIVE, age_days=10) -> ApprovedSharePointWorkbook:
    return ApprovedSharePointWorkbook(
        tenant_id=TENANT, site_id=site, drive_id=drive, item_id=item,
        etag=f"etag-{item}", web_url=f"https://aot.sharepoint.com/item/{item}",
        filename=name, modified_at=NOW - timedelta(days=age_days),
        content=workbook_bytes(sheet),
    )


def index(*documents: ApprovedSharePointWorkbook):
    return ApprovedCustomerEvidenceIndex.from_documents(
        documents=documents,
        tenant_id=TENANT,
        allowed_site_ids=frozenset({SITE}),
        allowed_drive_ids=frozenset({DRIVE}),
        now=NOW,
    )


def test_aot_delivery_receipt_matches_exact_serial_number():
    doc = source(
        name="AOT Delivery Receipt Equipment 2026.xlsx",
        sheet=[
            ("Receipt - 1", [
                ("", "CO. NAME:", "Marcella Operator LLC", ""),
                ("", "ID#:", "1403", ""),
                ("Copier", "Kyocera MZ6001i", "S/N:", "1G55503749"),
            ]),
        ],
    )
    result = index(doc)
    found = result.match(serial_number="1G55503749")
    assert found.status == "verified"
    assert found.customer_name == "Marcella Operator LLC"
    assert len(found.evidence) == 1
    assert found.evidence[0].source_item_id == doc.item_id
    assert found.evidence[0].source_etag == doc.etag
    assert len(found.evidence[0].source_sha256) == 64
    assert result.match(serial_number="1G5550374").status == "not_found"


def test_aot_copier_log_keeps_serials_scoped_to_their_own_sheet():
    doc = source(
        name="Virginia Health Services Copier Log.xlsx",
        sheet=[
            ("Marcella", [
                ("", "Company:", ""),
                ("", "Marcella Operator LLC", ""),
                ("", "MZ6001i", "1403"),
                ("", "Serial #:", "1G55503749"),
            ]),
            ("Atlantic", [
                ("", "Company:", ""),
                ("", "Atlantic View Operator", ""),
                ("", "MZ6001i", "1398"),
                ("", "Serial #:", "1G55504297"),
            ]),
            ("Water Reed", [
                ("", "Company:", ""),
                ("", "Water Reed Operator", ""),
                ("", "MZ6001i", "1405"),
                ("", "Serial #:", "1G55503694"),
            ]),
        ],
    )
    mapped = index(doc)
    assert mapped.match(serial_number="1G55503749").customer_name == "Marcella Operator LLC"
    assert mapped.match(serial_number="1G55504297").customer_name == "Atlantic View Operator"
    assert mapped.match(serial_number="1G55503694").customer_name == "Water Reed Operator"
    assert mapped.match(serial_number="RRG2301432").status == "not_found"


def test_distinct_documents_with_agreeing_name_are_verified():
    a = source(sheet=[("One", [("CO. NAME:", "Marcella Operator LLC"), ("SN:", "1G55503749")])])
    b = source(
        sheet=[("Two", [("CO. NAME:", "Marcella Operator LLC"), ("S/N:", "1G55503749")])],
        item="other-source",
    )
    match = index(a, b).match(serial_number="1G55503749")
    assert match.status == "verified"
    assert len(match.evidence) == 2


def test_conflicting_source_names_fail_closed():
    a = source(sheet=[("One", [("CO. NAME:", "Marcella Operator LLC"), ("SN:", "1G55503749")])])
    b = source(sheet=[("Two", [("CO. NAME:", "Other Customer"), ("SN:", "1G55503749")])], item="other")
    match = index(a, b).match(serial_number="1G55503749")
    assert match.status == "conflict"
    assert match.customer_name is None
    assert len(match.evidence) == 2


def test_disagreement_with_existing_kfs_customer_fails_closed():
    doc = source(sheet=[("One", [("CO. NAME:", "Marcella Operator LLC"), ("SN:", "1G55503749")])])
    match = index(doc).match(serial_number="1G55503749", kfs_customer="Different Company")
    assert match.status == "conflict"
    assert match.customer_name is None


@pytest.mark.parametrize("field,value", [
    ("site", "different-site"),
    ("drive", "different-library"),
])
def test_unknown_document_scope_is_rejected_before_parsing(field, value):
    opts = {field: value}
    doc = source(sheet=[("One", [("CO. NAME:", "Client"), ("SN:", "1G55503749")])], **opts)
    with pytest.raises(PermissionError, match="outside approved"):
        index(doc)


def test_missing_scope_is_rejected():
    doc = source(sheet=[("One", [("CO. NAME:", "Client"), ("SN:", "1G55503749")])])
    with pytest.raises(PermissionError, match="approved tenant"):
        ApprovedCustomerEvidenceIndex.from_documents(
            documents=(doc,), tenant_id=TENANT,
            allowed_site_ids=frozenset(), allowed_drive_ids=frozenset({DRIVE}), now=NOW,
        )


def test_stale_customer_source_is_ignored():
    doc = source(sheet=[("One", [("CO. NAME:", "Client"), ("SN:", "1G55503749")])], age_days=500)
    assert index(doc).match(serial_number="1G55503749").status == "not_found"


def test_duplicate_versioned_document_not_counted_twice():
    doc = source(sheet=[("One", [("CO. NAME:", "Client"), ("SN:", "1G55503749")])])
    with pytest.raises(ValueError, match="duplicate"):
        index(doc, doc)


def test_vendor_and_bare_serial_without_customer_label_do_not_create_mapping():
    doc = source(sheet=[("One", [("eCopier Solutions",), ("Serial #:", "1G55503749")])])
    assert index(doc).match(serial_number="1G55503749").status == "not_found"


def test_unauthorized_or_macro_enabled_file_is_rejected():
    doc = source(sheet=[("One", [("CO. NAME:", "Client"), ("SN:", "1G55503749")])])
    with pytest.raises(ValueError, match="non-macro"):
        ApprovedSharePointWorkbook(
            tenant_id=doc.tenant_id, site_id=doc.site_id, drive_id=doc.drive_id,
            item_id=doc.item_id, etag=doc.etag, web_url=doc.web_url,
            filename="Copier log.xlsm", modified_at=doc.modified_at, content=doc.content,
        )
